#!/usr/bin/env python3
"""Coach performance report (PRD_COACH_REPORTS.md) → .xlsx + Telegram summary.

Calls the analytics-coaches edge function (?action=report, plus
?action=cohort_matrix for month and longer periods) and writes:

  coach-report-<period>-<start>.xlsx          sheets: Ranking, Per-coach detail,
                                              Cohort churn (month+ only),
                                              Tournaments, Data notes
  coach-report-<period>-<start>-summary.txt   short Telegram text: school
                                              headline, top 3, bottom 3, movers

Prints {"xlsx": ..., "summary": ...} on stdout. Read-only.

Usage:
  python3 scripts/coach_report.py --period week|month|quarter|half|year \
      [--date YYYY-MM-DD] --out DIR
  --date is any day inside the wanted period; default = yesterday in Almaty
  (so a run on Monday / the 1st reports the period that just ended).

Environment (no secrets in this file):
  CE_SERVICE_KEY   service-role / sb_secret key for the Chess Empire project (required)
  CE_SUPABASE_URL  project URL (required)
"""

import argparse
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

PERIODS = ('week', 'month', 'quarter', 'half', 'year')
RAZRYADS = ('none', '4th', '3rd', '2nd', '1st', 'other')
BUCKETS = ('0', '1', '2', '3', '4+')
KINDS = ('league_c', 'league_b', 'razryad_4', 'razryad_3', 'rated')
METRICS = ('retention', 'net_growth', 'razryad_promotions', 'tournament_participation', 'podiums')

DATA_NOTES = [
    'Tournament and razryad metrics start in May 2026 (tournament_results from 2026-05-13, '
    'razryad_history from 2026-05-18), so the yearly 2026 report is partial.',
    'Retention for students who started before Dec 2025 is a lower bound (attendance starts 2025-12-01; '
    'their start is marked "start ≤ Dec 2025").',
    'The weekly cohort matrix is too thin, so the Cohort churn sheet only appears in the monthly report and longer ones.',
]

DEFINITIONS = [
    'Counts are credited to the coach at the time of the event (audit_log coach_id timeline, complete since '
    '2026-02-10); unknown or deleted coach → "Unattributed". Razryad changes use razryad_history.trainer_id_at_change.',
    'New = first activation in the period (new student row, or left → active after more than 60 days).',
    'Left and Frozen 30+ (frozen 30+ days and not reactivated by period end) are separate; churn = left + frozen 30+.',
    'Net growth = new − left. Duplicate status events (same transition within 5 minutes) are counted once.',
    'Retention: start (first present attendance, else created_at) → first leave; Kaplan–Meier survival at 3 / 6 months.',
    'Active roster, razryad snapshot and tournament participation buckets use the current roster (snapshot at run time).',
    'Small sample = fewer than 10 active students.',
]


# ---------------------------------------------------------------- fetching

def fetch(base_url, key, params):
    url = f"{base_url.rstrip('/')}/functions/v1/analytics-coaches?{urllib.parse.urlencode(params)}"
    req = urllib.request.Request(url, headers={'Authorization': f'Bearer {key}', 'apikey': key})
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = json.load(resp)
    except urllib.error.HTTPError as e:
        raise RuntimeError(f'analytics-coaches {params.get("action")}: HTTP {e.code} {e.read().decode(errors="replace")[:300]}')
    if not body.get('success'):
        raise RuntimeError(f'analytics-coaches {params.get("action")}: {body.get("error")}')
    return body['data']


def fetch_payload(base_url, key, period, date=None):
    params = {'action': 'report', 'period': period}
    if date:
        params['date'] = date
    report = fetch(base_url, key, params)
    cohort = fetch(base_url, key, {'action': 'cohort_matrix'}) if period != 'week' else None
    return {'report': report, 'cohort': cohort}


# ---------------------------------------------------------------- helpers

def pct(v):
    return None if v is None else round(float(v) * 100, 1)


def share(s):
    """{'count': n, 'pct': p} → 'n (p%)'."""
    if not s:
        return ''
    return f"{s.get('count', 0)} ({s['pct']}%)" if s.get('pct') is not None else str(s.get('count', 0))


def played_pct(row):
    active = row.get('active_students') or 0
    zero = (((row.get('participation') or {}).get('period') or {}).get('0') or {}).get('count', active)
    return round(100.0 * (active - zero) / active, 1) if active else None


def by_razryad(d):
    return ', '.join(f'{k}: {v}' for k, v in sorted((d or {}).items()))


def ranked(report):
    return [c for c in report['coaches'] if c.get('rank') is not None]


def extra_rows(report):
    return [r for r in (report.get('unattributed'), report.get('school')) if r]


def window_label(w):
    return f"{w['start_date']} … {w['end_date']}"


# ---------------------------------------------------------------- workbook

def _sheet(wb, title, header, rows, widths=None):
    from openpyxl.styles import Font
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'B2'
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return ws


def build_workbook(payload, path):
    from openpyxl import Workbook
    report, cohort = payload['report'], payload.get('cohort')
    prev = {c['coach_id']: c for c in report.get('previous', {}).get('coaches', [])}
    wb = Workbook()
    wb.remove(wb.active)

    # Ranking
    rows = []
    for c in report['coaches'] + extra_rows(report):
        r = c.get('retention') or {}
        p = prev.get(c.get('coach_id')) if c.get('coach_id') else None
        rows.append([
            c.get('rank'), c['coach_name'], c.get('score'), 'yes' if c.get('small_sample') else '',
            c.get('active_students'), c.get('new_students'), c.get('left'), c.get('frozen_30d'), c.get('churn'),
            c.get('net_growth'), pct(r.get('km_survival_3m')), pct(r.get('km_survival_6m')),
            (c.get('razryad_changes') or {}).get('promotions'), played_pct(c),
            (c.get('podiums') or {}).get('pct_of_participants'),
            p.get('rank') if p else None,
            round(c['score'] - p['score'], 1) if p and c.get('score') is not None and p.get('score') is not None else None,
        ] + [(c.get('ranks') or {}).get(m) for m in METRICS])
    _sheet(wb, 'Ranking', [
        'Rank', 'Coach', 'Score (0–100)', 'Small sample', 'Active', 'New', 'Left', 'Frozen 30+', 'Churn',
        'Net growth', 'Survival 3m %', 'Survival 6m %', 'Razryad promotions', 'Played ≥1 tournament %',
        'Podium % of participants', 'Previous rank', 'Score Δ',
        'Rank: retention', 'Rank: net growth', 'Rank: promotions', 'Rank: participation', 'Rank: podiums',
    ], rows, [7, 28])

    # Per-coach detail
    rows = []
    for c in report['coaches'] + extra_rows(report):
        r = c.get('retention') or {}
        snap = c.get('razryad_snapshot') or {}
        rz = c.get('razryad_changes') or {}
        rows.append([
            c['coach_name'], c.get('active_students'), c.get('new_students'), c.get('left'), c.get('frozen_30d'),
            c.get('churn'), c.get('net_growth'), r.get('n'), r.get('events'), r.get('left_censored'),
            r.get('mean_months'), r.get('median_months'), pct(r.get('km_survival_3m')), pct(r.get('km_survival_6m')),
        ] + [share(snap.get(k)) for k in RAZRYADS] + [
            rz.get('promotions'), rz.get('downgrades'),
            by_razryad(rz.get('promotions_by_razryad')), by_razryad(rz.get('downgrades_by_razryad')),
        ])
    _sheet(wb, 'Per-coach detail', [
        'Coach', 'Active', 'New', 'Left', 'Frozen 30+', 'Churn', 'Net growth',
        'Retention: students', 'Retention: left', 'Retention: start ≤ Dec 2025', 'Mean months', 'Median months',
        'Survival 3m %', 'Survival 6m %',
    ] + [f'Razryad {k}' for k in RAZRYADS] + [
        'Promotions', 'Downgrades', 'Promotions by razryad', 'Downgrades by razryad',
    ], rows, [28])

    # Cohort churn (month and longer)
    if cohort is not None:
        n = cohort.get('max_months', 0) + 1
        header = ['Start month', 'Cohort size', 'Start ≤ Dec 2025']
        rows = []
        for label, key in (('Left %', 'left_pct'), ('Left + frozen 30+ %', 'churn_pct')):
            rows.append([label])
            for r in cohort.get('rows', []):
                rows.append([r['start_month'], r['cohort_size'], 'yes' if r.get('left_censored') else ''] + list(r.get(key) or []))
            rows.append([])
        rows.append(['Exits by level at leaving', 'Left', 'Left + frozen 30+'])
        rows += [[f"Level {x['level']}", x['left'], x['churn']] for x in cohort.get('by_level_at_leaving', [])]
        rows.append([])
        rows.append(['Exits by razryad at leaving', 'Left', 'Left + frozen 30+'])
        rows += [[x['razryad'], x['left'], x['churn']] for x in cohort.get('by_razryad_at_leaving', [])]
        _sheet(wb, 'Cohort churn', header + [f'M{k}' for k in range(n)], rows, [28, 12, 16])

    # Tournaments
    rows = []
    for c in report['coaches'] + extra_rows(report):
        part = c.get('participation') or {}
        pod = c.get('podiums') or {}
        rows.append([c['coach_name'], c.get('active_students')]
                    + [share((part.get('period') or {}).get(b)) for b in BUCKETS]
                    + [share((part.get('lifetime') or {}).get(b)) for b in BUCKETS]
                    + [pod.get('participants'), pod.get('students'), pod.get('finishes'), pod.get('pct_of_participants')]
                    + [(pod.get('by_kind') or {}).get(k) for k in KINDS])
    _sheet(wb, 'Tournaments', ['Coach', 'Active']
           + [f'Period: {b} tournaments' for b in BUCKETS]
           + [f'Lifetime: {b} tournaments' for b in BUCKETS]
           + ['Participants (period)', 'Podium students', 'Podium finishes', 'Podium % of participants']
           + [f'Podiums {k}' for k in KINDS], rows, [28])

    # Data notes
    w = report['period']
    rows = [
        ['Period', f"{w['period']} {window_label(w)} (Asia/Almaty)"],
        ['Previous period', window_label(report['previous_period'])],
        ['Generated at', report.get('generated_at') or ''],
        ['Score weights', ', '.join(f'{k} {v}' for k, v in (report.get('weights') or {}).items())],
        [],
        ['Data limits'],
    ] + [['', note] for note in DATA_NOTES] + [[], ['Definitions']] + [['', d] for d in DEFINITIONS]
    ws = _sheet(wb, 'Data notes', ['Item', 'Value'], rows, [18, 120])
    ws.freeze_panes = None

    wb.save(path)
    return path


# ---------------------------------------------------------------- summary

def _line(c):
    tag = ' (small sample)' if c.get('small_sample') else ''
    return f"{c['rank']}. {c['coach_name']} — {c['score']}{tag}"


def build_summary(payload):
    report = payload['report']
    w = report['period']
    s = report.get('school') or {}
    rz = s.get('razryad_changes') or {}
    pod = s.get('podiums') or {}
    rt = s.get('retention') or {}
    coaches = ranked(report)
    top = coaches[:3]
    bottom = [c for c in coaches[-3:] if c not in top]

    lines = [
        f"Coach report — {w['period']} {window_label(w)}",
        f"School: active {s.get('active_students', 0)} · new {s.get('new_students', 0)} · left {s.get('left', 0)}"
        f" · frozen 30+ {s.get('frozen_30d', 0)} · churn {s.get('churn', 0)} · net {s.get('net_growth', 0):+d}",
        f"Razryads: +{rz.get('promotions', 0)} promotions, {rz.get('downgrades', 0)} downgrades"
        f" · Podiums: {pod.get('students', 0)} students, {pod.get('finishes', 0)} finishes"
        f" · Survival 6m: {pct(rt.get('km_survival_6m')) if rt.get('km_survival_6m') is not None else 'n/a'}%",
        '',
        'Top 3:',
    ] + [_line(c) for c in top] + ['', 'Bottom 3:'] + [_line(c) for c in bottom]

    mv = report.get('movers') or {}
    moves = [('▲', m) for m in mv.get('up', [])] + [('▼', m) for m in mv.get('down', [])]
    if moves:
        lines += ['', 'Movers vs previous period:']
        lines += [f"{a} {m['coach_name']} {m['score_delta']:+} (rank {m['previous_rank']} → {m['rank']})" for a, m in moves]
    return '\n'.join(lines) + '\n'


def write_report(payload, out_dir):
    report = payload['report']
    w = report['period']
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.join(out_dir, f"coach-report-{w['period']}-{w['start_date']}")
    xlsx = build_workbook(payload, base + '.xlsx')
    summary = base + '-summary.txt'
    with open(summary, 'w', encoding='utf-8') as f:
        f.write(build_summary(payload))
    return {'xlsx': xlsx, 'summary': summary}


def main(argv=None):
    ap = argparse.ArgumentParser(description='Coach performance report → xlsx + Telegram summary')
    ap.add_argument('--period', required=True, choices=PERIODS)
    ap.add_argument('--date', help='YYYY-MM-DD inside the period (default: yesterday in Almaty)')
    ap.add_argument('--out', required=True, help='output directory')
    args = ap.parse_args(argv)

    key = os.environ.get('CE_SERVICE_KEY')
    base_url = os.environ.get('CE_SUPABASE_URL')
    if not key or not base_url:
        print('CE_SERVICE_KEY and CE_SUPABASE_URL must be set', file=sys.stderr)
        return 2
    try:
        payload = fetch_payload(base_url, key, args.period, args.date)
    except (RuntimeError, urllib.error.URLError) as e:
        print(str(e), file=sys.stderr)
        return 1
    print(json.dumps(write_report(payload, args.out), ensure_ascii=False))
    return 0


if __name__ == '__main__':
    sys.exit(main())
